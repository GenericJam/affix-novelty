"""Build English single-affix entries from MorphoLEX validated segmentation."""
import re
import openpyxl
from wordfreq import word_frequency, zipf_frequency

WORD_FREQ_FLOOR = 1.5


def _parse(seg):
    return (re.findall(r"<([a-z]+)<", seg),
            re.findall(r"\(([a-z]+)\)", seg),
            re.findall(r">([a-z]+)>", seg))


def load(path="MorphoLEX_en.xlsx"):
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = [s for s in wb.sheetnames if re.fullmatch(r"\d+-\d+-\d+", s)]
    raw = []
    for sh in sheets:
        it = wb[sh].iter_rows(values_only=True)
        header = list(next(it))
        try:
            wi = header.index("Word"); si = header.index("MorphoLexSegm")
        except ValueError:
            continue
        for r in it:
            word, seg = r[wi], r[si]
            if not word or not seg:
                continue
            w = str(word).lower()
            if not w.isalpha():
                continue
            p, ro, s = _parse(seg)
            raw.append((w, p, ro, s))

    def keep(p, ro, s, w):
        return len(ro) == 1 and zipf_frequency(w, "en") >= WORD_FREQ_FLOOR

    pref = [(w, p[0], ro[0]) for w, p, ro, s in raw
            if len(p) == 1 and len(s) == 0 and keep(p, ro, s, w)]
    suff = [(w, s[0], ro[0]) for w, p, ro, s in raw
            if len(p) == 0 and len(s) == 1 and keep(p, ro, s, w)]

    def make(items, side):
        # dedup inflections -> shortest (lemma), tie-break frequency
        best = {}
        for w, af, root in items:
            key = (af, root)
            rank = (len(w), -word_frequency(w, "en"))
            if key not in best or rank < best[key][0]:
                best[key] = (rank, w)
        out = []
        for (af, root), (_, w) in best.items():
            out.append(dict(word=w, side=side, affix=af, root=root,
                            freq=word_frequency(w, "en"), zipf=zipf_frequency(w, "en")))
        return out

    return make(pref, "prefix") + make(suff, "suffix")
