"""
Prefixes vs suffixes as generators of non-compositional ("new meaning") words,
using VALIDATED morphology from MorphoLEX (Sanchez-Gutierrez et al. 2018) instead
of orthographic guessing.

Two upgrades over the orthographic version:
  1. Bound-root prefix words (receive, conceive, reduce, retain) are now in scope.
     They have no standalone root vector, so we reconstruct the root's meaning from
     its morphological family (each sibling minus its affix offset, leave-one-out)
     and score the word against affix_offset + reconstructed_root.
  2. Token-level (usage-frequency-weighted) novelty alongside type-level.

Compositionality (high = predictable from parts, low = drifted/novel):
     comp(word) = cosine( vec(word),  affix_offset(affix) + R )
where R is vec(root) for a free singleton root, else the family reconstruction.
Affix offsets are learned (two-pass, self-cleaning) from that affix's free-root words.
"""

import re
import csv
import numpy as np
import openpyxl
import gensim.downloader as api
from wordfreq import word_frequency, zipf_frequency
from scipy.stats import mannwhitneyu

THRESHOLDS = [0.10, 0.15, 0.20]
MIN_OFFSET_PAIRS = 8        # affix needs this many free-root words to learn an offset
WORD_FREQ_FLOOR  = 1.5      # zipf floor on the derived word (keep real words, allow rare)

# ----------------------------------------------------------------- load vectors
print("loading GloVe vectors...")
KV = api.load("glove-wiki-gigaword-300")
def vec(w): return KV[w] if w in KV else None
def cos(a, b): return float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b)))

# ----------------------------------------------------------------- parse MorphoLEX
print("parsing MorphoLEX...")
wb = openpyxl.load_workbook("MorphoLEX_en.xlsx", read_only=True)
PRS_SHEETS = [s for s in wb.sheetnames if re.fullmatch(r"\d+-\d+-\d+", s)]

def parse_seg(seg):
    return (re.findall(r"<([a-z]+)<", seg),     # prefixes
            re.findall(r"\(([a-z]+)\)", seg),   # roots
            re.findall(r">([a-z]+)>", seg))     # suffixes

entries = []
for sh in PRS_SHEETS:
    ws = wb[sh]
    it = ws.iter_rows(values_only=True)
    header = [h for h in next(it)]
    # header row sometimes shifted; find columns by name within first cells
    try:
        wi = header.index("Word"); si = header.index("MorphoLexSegm")
    except ValueError:
        continue
    for r in it:
        word, seg = r[wi], r[si]
        if not word or not seg:
            continue
        w = str(word).lower()
        if not w.isalpha():
            continue
        prefs, roots, suffs = parse_seg(seg)
        entries.append(dict(word=w, prefs=prefs, roots=roots, suffs=suffs))

# core contrast: exactly one affix, exactly one root, word has a vector & is real
def keep(e):
    return (len(e["roots"]) == 1 and vec(e["word"]) is not None
            and zipf_frequency(e["word"], "en") >= WORD_FREQ_FLOOR)

PREF = [e for e in entries if keep(e) and len(e["prefs"]) == 1 and len(e["suffs"]) == 0]
SUFF = [e for e in entries if keep(e) and len(e["suffs"]) == 1 and len(e["prefs"]) == 0]
for e in PREF: e.update(side="prefix", affix=e["prefs"][0], root=e["roots"][0])
for e in SUFF: e.update(side="suffix", affix=e["suffs"][0], root=e["roots"][0])

def dedup(group):
    """Collapse inflectional variants of one lexeme (same affix+root) to a single
    representative: the SHORTEST surface form (the lemma), tie-broken by frequency.
    So receive/received/receiving -> receive, giving the cleanest vector and label."""
    best = {}
    for e in group:
        key = (e["side"], e["affix"], e["root"])
        rank = (len(e["word"]), -word_frequency(e["word"], "en"))  # shorter, then commoner
        if key not in best or rank < best[key][0]:
            best[key] = (rank, e)
    return [e for _, e in best.values()]

PREF, SUFF = dedup(PREF), dedup(SUFF)
ALL = PREF + SUFF
for e in ALL:
    e["freq"] = word_frequency(e["word"], "en")
    e["zipf"] = zipf_frequency(e["word"], "en")
    e["root_free"] = vec(e["root"]) is not None
print(f"  {len(PREF)} prefix lexemes, {len(SUFF)} suffix lexemes "
      f"(after dedup; bound-root: {sum(not e['root_free'] for e in ALL)})")

# ----------------------------------------------------------------- learn offsets
def learn_offsets(items):
    by = {}
    for e in items:
        if e["root_free"]:
            by.setdefault((e["side"], e["affix"]), []).append(e)
    offs = {}
    for key, grp in by.items():
        diffs = [vec(e["word"]) - vec(e["root"]) for e in grp]
        if len(diffs) < MIN_OFFSET_PAIRS:
            continue
        o1 = np.mean(diffs, axis=0)
        comps = [cos(vec(e["word"]), vec(e["root"]) + o1) for e in grp]
        med = np.median(comps)
        keep_d = [d for d, c in zip(diffs, comps) if c >= med]
        offs[key] = np.mean(keep_d, axis=0) if keep_d else o1
    return offs

OFFSETS = learn_offsets(ALL)
print(f"  learned offsets for {len(OFFSETS)} affixes")

# ----------------------------------------------------------------- families
# family of a root = all core lexemes sharing that root (any affix, either side)
FAMILY = {}
for e in ALL:
    FAMILY.setdefault(e["root"], []).append(e)

def reconstruct_root(root, exclude_word):
    """Leave-one-out estimate of the root's meaning from its siblings."""
    hats = []
    for s in FAMILY[root]:
        if s["word"] == exclude_word:
            continue
        key = (s["side"], s["affix"])
        if key in OFFSETS:
            hats.append(vec(s["word"]) - OFFSETS[key])
    return np.mean(hats, axis=0) if hats else None

# ----------------------------------------------------------------- score
scored = []
for e in ALL:
    key = (e["side"], e["affix"])
    if key not in OFFSETS:
        continue
    R = reconstruct_root(e["root"], e["word"])      # family reconstruction (LOO)
    method = "family"
    if R is None:                                   # no usable siblings
        if not e["root_free"]:
            continue                                # bound singleton: cannot score
        R = vec(e["root"]); method = "freeroot"
    e["comp"] = cos(vec(e["word"]), OFFSETS[key] + R)
    e["method"] = method
    scored.append(e)

SPREF = [e for e in scored if e["side"] == "prefix"]
SSUFF = [e for e in scored if e["side"] == "suffix"]

# ----------------------------------------------------------------- calibration
print("\n--- calibration (low comp = novel/drifted) ---")
bw = {e["word"]: e for e in scored}
def show(w):
    e = bw.get(w)
    if e:
        print(f"  {w:12s} {e['side'][:4]} {e['affix']:6s}+{e['root']:9s} "
              f"comp={e['comp']:+.3f} [{e['method']},{'free' if e['root_free'] else 'BOUND'}]")
    else:
        print(f"  {w:12s} -- not scored")
for w in ["report","reduce","retain","rebuild","reread","display","predict",
          "business","happiness","treatment","fellowship","payment","careless"]:
    show(w)
print("  -- bound-root families (should mostly read as NOVEL) --")
for root in ["ceive","duct","sume","tain","cur","fer","mit"]:
    fam = sorted([e for e in scored if e["root"] == root], key=lambda e: e["comp"])
    if fam:
        print(f"  ({root:5s}) " + ", ".join(f"{e['word']}({e['comp']:+.2f})" for e in fam))

# ----------------------------------------------------------------- aggregate
def rates(items, weighted=False):
    c = np.array([e["comp"] for e in items])
    w = np.array([e["freq"] for e in items]) if weighted else np.ones(len(items))
    w = w / w.sum()
    out = {f"<{t}": float(w[c < t].sum()) for t in THRESHOLDS}
    out["mean"] = float(np.average(c, weights=w)); out["n"] = len(items)
    return out

print("\n--- prefix vs suffix : TYPE level (each lexeme once) ---")
for name, items in (("prefix", SPREF), ("suffix", SSUFF)):
    r = rates(items)
    print(f"  {name:7s} n={r['n']:5d} mean={r['mean']:+.3f}  " +
          "  ".join(f"novel{t}={r[t]:.1%}" for t in [f'<{x}' for x in THRESHOLDS]))
U, p = mannwhitneyu([e["comp"] for e in SPREF], [e["comp"] for e in SSUFF],
                    alternative="two-sided")
print(f"  Mann-Whitney p={p:.2e}")

print("\n--- prefix vs suffix : TOKEN level (weighted by usage frequency) ---")
for name, items in (("prefix", SPREF), ("suffix", SSUFF)):
    r = rates(items, weighted=True)
    print(f"  {name:7s}        mean={r['mean']:+.3f}  " +
          "  ".join(f"novel{t}={r[t]:.1%}" for t in [f'<{x}' for x in THRESHOLDS]))

print("\n--- bound-root subset only (the words orthography missed) ---")
for name, items in (("prefix", [e for e in SPREF if not e["root_free"]]),
                    ("suffix", [e for e in SSUFF if not e["root_free"]])):
    if items:
        r = rates(items)
        print(f"  {name:7s} n={r['n']:5d} mean={r['mean']:+.3f}  novel<0.15={r['<0.15']:.1%}")

print("\n--- novelty rate (comp<0.15) controlled for word frequency ---")
print(f"  {'zipf':12s}{'pref n':>8s}{'pref nov':>10s}{'suf n':>8s}{'suf nov':>10s}")
for lo, hi in [(1.5,2.5),(2.5,3.0),(3.0,3.5),(3.5,4.0),(4.0,7.0)]:
    pr = [e for e in SPREF if lo <= e["zipf"] < hi]
    su = [e for e in SSUFF if lo <= e["zipf"] < hi]
    prr = np.mean([e["comp"]<0.15 for e in pr]) if pr else float("nan")
    sur = np.mean([e["comp"]<0.15 for e in su]) if su else float("nan")
    print(f"  [{lo:.1f},{hi:.1f}){'':3s}{len(pr):8d}{prr:10.1%}{len(su):8d}{sur:10.1%}")

# ----------------------------------------------------------------- per affix
def affix_table(items, side, top=12):
    by = {}
    for e in items:
        by.setdefault(e["affix"], []).append(e)
    rows = []
    for af, grp in by.items():
        c = np.array([e["comp"] for e in grp])
        rows.append((af, len(grp), c.mean(), float((c < 0.15).mean()), grp))
    rows = [r for r in rows if r[1] >= 12]
    rows.sort(key=lambda x: x[3], reverse=True)
    print(f"\n--- most word-generating {side}es (novel rate, n>=12) ---")
    for af, n, m, rate, grp in rows[:top]:
        ex = ", ".join(f"{e['word']}({e['comp']:+.2f})"
                       for e in sorted(grp, key=lambda e: e["comp"])[:4])
        print(f"  {af:8s} n={n:4d} novel={rate:5.1%} mean={m:+.2f}  e.g. {ex}")

affix_table(SPREF, "prefix")
affix_table(SSUFF, "suffix")

# ----------------------------------------------------------------- dump
with open("results_morpholex.csv", "w", newline="") as f:
    wr = csv.DictWriter(f, fieldnames=["word","side","affix","root","root_free",
                                       "method","zipf","freq","comp"])
    wr.writeheader()
    for e in sorted(scored, key=lambda e: e["comp"]):
        wr.writerow({k: e[k] for k in wr.fieldnames})
print(f"\nwrote {len(scored)} scored lexemes to results_morpholex.csv")
